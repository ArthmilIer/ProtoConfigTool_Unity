import os
from openpyxl import load_workbook
from tools.field_parser import parse_fields

def row_is_empty(row):
    return all(v is None or v == "" for v in row)

def trim_trailing_empty_rows(rows):
    end = len(rows)
    while end > 0 and row_is_empty(rows[end - 1]):
        end -= 1
    return rows[:end]

def load_excels(excel_dir):
    result = []

    for file in os.listdir(excel_dir):
        if not file.endswith(".xlsx") or file.startswith("~$"):
            continue

        path = os.path.join(excel_dir, file)
        wb = load_workbook(path, data_only=True)

        excel_info = {
            "name": os.path.splitext(file)[0],
            "sheets": []
        }

        for sheet in wb.worksheets:
            fields = parse_fields(sheet[1])
            data = list(sheet.iter_rows(min_row=3, values_only=True))
            data = trim_trailing_empty_rows(data)

            excel_info["sheets"].append({
                "excel": excel_info["name"],
                "sheet": sheet.title,
                "fields": fields,
                "rows": data
            })

        result.append(excel_info)

    return result
